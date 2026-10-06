#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Tests for the WOD burnt-area and cause statistics applications.

The fires are inserted through the ORM: what has to be asserted is arithmetic over
known areas in known years, and building that by hand is clearer than arranging for
the importer to produce it. The fixture spans two calendar years and four fire years,
both agencies, all four cause classes, two perimeters and one point in the sea, so
that every column, both year conventions, both area methods and both country sources
have something to tell apart.
"""

import csv
import datetime
import logging

import pytest

from shapely.geometry import MultiPolygon
from shapely.geometry import box

from src.apps.statistics.wildfires.new_zealand_wod import wildfire_causes as causes
from src.apps.statistics.wildfires.new_zealand_wod import wildfire_statistics as app
from src.data_model.data_provider import DataProvider
from src.providers import new_zealand_wod
from src.providers import ocha
from src.providers.new_zealand_wod.ignition import WodIgnition
from src.providers.new_zealand_wod.wildfire import WodWildfire
from src.providers.ocha.admin_boundary import OchaAdminBoundary

logger = logging.getLogger("test-wod-statistics")

UTC = datetime.timezone.utc

NEW_ZEALAND = box(166.0, -48.0, 179.0, -34.0)

#: A square of side ``metres`` at an NZTM origin, as a MULTIPOLYGON.
def square(x: float, y: float, metres: float) -> str:
    return f"SRID=2193;{MultiPolygon([box(x, y, x + metres, y + metres)]).wkt}"


#: (wod_id, source, date, cause, area_ha, (lon, lat), perimeter).
FIRES = [
    ("WOD1", "FENZ", datetime.date(2020, 1, 10), "lightning", 50.0, (172.0, -43.0),
     square(1570000, 5180000, 1000)),                 # 100 ha measured
    ("WOD2", "FENZ", datetime.date(2020, 8, 10), "human", 0.5, (174.0, -41.0), None),
    ("WOD3", "FENZ", datetime.date(2020, 3, 1), "unknown", 2.0, (175.0, -39.0), None),
    ("WOD4", "DOC", datetime.date(2005, 2, 1), "natural", 10.0, (176.0, -38.0),
     square(1580000, 5190000, 200)),                  # 4 ha measured
    ("WOD5", "DOC", datetime.date(2005, 11, 1), "human", 0.0, (171.0, -44.0), None),
    # In the sea, east of North Cape: a New Zealand fire recorded at a wrong location.
    ("WOD6", "FENZ", datetime.date(2020, 5, 1), "human", 1.0, (179.5, -34.5), None),
]


def season(month: int) -> str:
    return {12: "Summer", 1: "Summer", 2: "Summer", 3: "Autumn", 4: "Autumn",
            5: "Autumn", 6: "Winter", 7: "Winter", 8: "Winter"}.get(month, "Spring")


@pytest.fixture
def populated(db_session):
    ocha_provider = DataProvider(name=ocha.PROVIDER_NAME, product=ocha.PROVIDER_PRODUCT,
                                 full_name=ocha.PROVIDER_FULL_NAME, url=ocha.PROVIDER_URL)
    wod_provider = DataProvider(name=new_zealand_wod.PROVIDER_NAME,
                                product=new_zealand_wod.PROVIDER_PRODUCT,
                                full_name=new_zealand_wod.PROVIDER_FULL_NAME,
                                url=new_zealand_wod.PROVIDER_URL)
    db_session.add_all([ocha_provider, wod_provider])
    db_session.flush()
    db_session.add(OchaAdminBoundary(
        data_provider_id=ocha_provider.id, source_id="NZL", level=0, name="New Zealand",
        geometry=f"SRID=4326;{MultiPolygon([NEW_ZEALAND]).wkt}",
        source="NZL", iso_code=554, iso_2="NZ", iso_3="NZL", iso_name="New Zealand",
        iso_3_group="NZL", region1_code=1, region1_name="r1", region2_code=2,
        region2_name="r2", region3_code=3, region3_name="r3", status_code=1,
        status_name="State", valid_date=datetime.date(2025, 1, 1),
        update_date=datetime.date(2025, 1, 1), land_source="osm", view="intl",
    ))

    for index, (wod_id, source, date, cause, area, point, perimeter) in enumerate(FIRES):
        instant = datetime.datetime(date.year, date.month, date.day, tzinfo=UTC)
        longitude, latitude = point
        ignition = WodIgnition(
            data_provider_id=wod_provider.id, wod_id=wod_id, year=date.year, source=source,
            geometry=f"SRID=4326;POINT({longitude} {latitude})",
            geometry_nztm=f"SRID=2193;POINT({1500000 + index} {5000000 + index})",
            date_time=instant, time_zone=new_zealand_wod.DEFAULT_TIME_ZONE)
        db_session.add(WodWildfire(
            data_provider_id=wod_provider.id, ignition=ignition, wod_id=wod_id,
            source=source, year=date.year, report_date=date, season=season(date.month),
            financial_year=date.year if date.month >= 7 else date.year - 1,
            area_ha=area, fire_cause=cause, perimeter_nztm=perimeter,
            start_date_time=instant, time_zone=new_zealand_wod.DEFAULT_TIME_ZONE))
    db_session.commit()
    return db_session


def run(session, **kwargs):
    return app.compute(session, kwargs.pop("year", None), logger, **kwargs)


def run_causes(session, **kwargs):
    return causes.compute(session, kwargs.pop("year", None), logger, **kwargs)


def find(rows, year, country="New Zealand"):
    matches = [row for row in rows if row.country == country and row.year == year]
    assert len(matches) == 1, f"expected one row for {country}/{year}, got {len(matches)}"
    return matches[0]


# --------------------------------------------------------------------------
# Burnt area
# --------------------------------------------------------------------------

def test_a_year_reports_its_fires_and_their_areas(populated):
    row = find(run(populated), 2020)
    assert (row.fires, row.minimum, row.maximum, row.total) == (4, 0.5, 50.0, 53.5)


def test_the_agency_and_perimeter_columns(populated):
    rows = run(populated)
    assert (find(rows, 2020).fenz, find(rows, 2020).doc, find(rows, 2020).perimeters) == (
        4, 0, 1)
    assert (find(rows, 2005).fenz, find(rows, 2005).doc, find(rows, 2005).perimeters) == (
        0, 2, 1)


def test_a_recorded_zero_is_counted(populated):
    row = find(run(populated), 2005)
    assert (row.fires, row.minimum) == (2, 0.0)


def test_the_years_come_newest_first_then_the_total(populated):
    rows = run(populated)
    assert [row.year_label for row in rows] == ["2020", "2005", "Total"]
    total = rows[-1]
    assert (total.fires, total.total, total.minimum, total.maximum) == (6, 63.5, 0.0, 50.0)
    assert (total.fenz, total.doc, total.perimeters) == (4, 2, 2)


def test_fire_years_run_from_july_to_june(populated):
    rows = run(populated, financial_year=True)
    assert [row.year_label for row in rows] == ["2020/2021", "2019/2020", "2005/2006",
                                                "2004/2005", "Total"]
    assert find(rows, 2019).fires == 3          # January, March and May 2020
    assert find(rows, 2020).fires == 1          # August 2020


def test_a_year_names_the_start_of_a_fire_year(populated):
    rows = run(populated, year=2019, financial_year=True)
    assert [row.year_label for row in rows] == ["2019/2020", "Total"]


def test_the_perimeter_method_measures_only_the_fires_with_one(populated):
    rows = run(populated, area_method=app.AREA_PERIMETER)
    assert find(rows, 2020).fires == 1
    assert find(rows, 2020).total == pytest.approx(100.0)
    assert find(rows, 2005).total == pytest.approx(4.0)


def test_a_minimum_area_selects_fires(populated):
    rows = run(populated, min_area=1.0)
    assert find(rows, 2020).fires == 3
    assert find(rows, 2005).fires == 1


def test_one_cause_and_one_agency_can_be_selected(populated):
    assert find(run(populated, cause="lightning"), 2020).fires == 1
    rows = run(populated, source="DOC")
    assert [row.year_label for row in rows] == ["2005", "Total"]


def test_filed_counts_the_point_in_the_sea(populated):
    assert find(run(populated), 2020).fires == 4


def test_geometry_leaves_the_point_in_the_sea_out(populated, caplog):
    with caplog.at_level(logging.WARNING):
        rows = run(populated, country_source=app.COUNTRY_SOURCE_GEOMETRY)
    assert find(rows, 2020).fires == 3
    assert "1 fire(s) have a point inside no country" in caplog.text


def test_unknown_vocabulary_is_refused():
    with pytest.raises(ValueError):
        app.scope_conditions(cause="Lightning discharge")
    with pytest.raises(ValueError):
        app.scope_conditions(source="Scion")
    with pytest.raises(ValueError):
        app.area_column("measured")
    with pytest.raises(ValueError):
        app.country_columns("everywhere")


def test_an_empty_year_reports_nothing(populated):
    assert run(populated, year=1950) == []


def test_the_csv_has_the_shared_columns_first(tmp_path, populated):
    path = tmp_path / "burnt.csv"
    app.write_csv(run(populated), path, logger)
    with path.open(encoding="utf-8") as handle:
        rows = list(csv.reader(handle))
    assert tuple(rows[0][:6]) == app.SHARED_COLUMNS
    assert rows[0][6:] == ["FENZ", "DOC", "Perimeters"]
    assert rows[1][:3] == ["New Zealand", "2020", "4"]


def test_the_docx_is_written_with_its_caveats(tmp_path, populated):
    from docx import Document

    path = tmp_path / "burnt.docx"
    app.write_docx(run(populated), path, logger, None, False, app.AREA_REPORTED, None,
                   app.COUNTRY_SOURCE_FILED, None, None)
    text = "\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs)
    assert "DOC records only" in text
    assert "five square metres" in text


def test_the_command_line(tmp_path):
    args = app.parse_arguments(["--csv", str(tmp_path / "x.csv")])
    assert (args.area_method, args.country_source, args.financial_year) == (
        "reported", "filed", False)
    with pytest.raises(SystemExit):
        app.parse_arguments([])
    with pytest.raises(SystemExit):
        app.parse_arguments(["--csv", "x.csv", "--min-area", "-1"])


# --------------------------------------------------------------------------
# Causes
# --------------------------------------------------------------------------

def test_a_year_counts_each_cause(populated):
    row = find(run_causes(populated), 2020)
    assert (row.fires, row.lightning, row.natural, row.human, row.unknown) == (4, 1, 0, 2, 1)
    assert row.fires == row.lightning + row.natural + row.human + row.unknown


def test_the_percentage_is_of_the_determined_fires(populated):
    row = find(run_causes(populated), 2020)
    assert row.determined == 3
    assert row.share == pytest.approx(100 / 3)


def test_the_area_share_is_of_the_determined_hectares(populated):
    """50 ha of lightning over 50 + 0.5 + 1 determined; the unknown 2 ha is in neither."""
    row = find(run_causes(populated), 2020)
    assert row.lightning_hectares == 50.0
    assert row.area_share == pytest.approx(100 * 50 / 51.5)


def test_natural_is_not_lightning(populated):
    row = find(run_causes(populated), 2005)
    assert (row.lightning, row.natural, row.share) == (0, 1, 0.0)


def test_the_total_shares_are_ratios_of_the_totals(populated):
    total = run_causes(populated)[-1]
    assert total.is_total
    assert (total.fires, total.determined, total.lightning) == (6, 5, 1)
    assert total.share == pytest.approx(20.0)


def test_a_share_of_nothing_is_no_answer():
    row = causes.Row("New Zealand", 2000, 1, 0, 0, 0, 1, 0.0, 0.0)
    assert row.share is None and row.area_share is None
    assert row.values[5] == "" and row.values[10] == ""


def test_the_fires_column_matches_the_companion_report(populated):
    for options in ({}, {"financial_year": True}, {"source": "FENZ"},
                    {"country_source": app.COUNTRY_SOURCE_GEOMETRY}):
        burnt = [(row.year, row.fires) for row in run(populated, **options)]
        counted = [(row.year, row.fires) for row in run_causes(populated, **options)]
        assert burnt == counted, options


def test_the_causes_csv(tmp_path, populated):
    path = tmp_path / "causes.csv"
    causes.write_csv(run_causes(populated), path, logger)
    with path.open(encoding="utf-8") as handle:
        rows = list(csv.reader(handle))
    assert tuple(rows[0]) == causes.COLUMNS
    assert rows[1][:9] == ["New Zealand", "2020", "4", "3", "1", "33.33", "0", "2", "1"]


def test_the_causes_docx_says_what_lightning_is(tmp_path, populated):
    from docx import Document

    path = tmp_path / "causes.docx"
    causes.write_docx(run_causes(populated), path, logger, None, False,
                      app.COUNTRY_SOURCE_FILED, None)
    text = "\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs)
    assert "Lightening Strike" in text
    assert "geothermal or volcanic" in text


def test_the_causes_command_line_has_no_cause_option(tmp_path):
    with pytest.raises(SystemExit):
        causes.parse_arguments(["--csv", "x.csv", "--cause", "lightning"])
    args = causes.parse_arguments(["--csv", str(tmp_path / "x.csv"), "--financial-year"])
    assert args.financial_year
