#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Tests for the CSIRO burnt-area statistics application.

The fires are inserted through the ORM: what has to be asserted is arithmetic over
known areas in known seasons. The measured areas are checked against :mod:`pyproj`,
an implementation independent of PostGIS.

The fixture carries what this report has that the others do not: **fire seasons**
that straddle the calendar year, a fire dated only by its extinction, **three fire
types** that have to add up, a **natural cause** published by some states and not
others, a fire before the default first season and one whose dates cannot be
believed.
"""

import argparse
import csv
import datetime
import logging

import pytest

from pyproj import Geod
from shapely.geometry import MultiPolygon
from shapely.geometry import box

from src.apps.statistics.wildfires.australia_csiro import wildfire_statistics as app
from src.data_model.data_provider import DataProvider
from src.providers import australia_csiro
from src.providers.australia_csiro.wildfire import CsiroWildfire

logger = logging.getLogger("test-csiro-statistics")

UTC = datetime.timezone.utc

GEOD = Geod(ellps="WGS84")

BUSHFIRE = australia_csiro.FIRE_TYPE_BUSHFIRE
PRESCRIBED = australia_csiro.FIRE_TYPE_PRESCRIBED_BURN
UNKNOWN = australia_csiro.FIRE_TYPE_UNKNOWN
NATURAL = australia_csiro.CAUSE_NATURAL
UNDETERMINED = australia_csiro.CAUSE_UNDETERMINED

#: (object_id, state, type, cause, ignition, extinguish, plausible, perimeter).
FIRES = [
    # 2019-20: two NSW fires either side of New Year, and a QLD one with no cause.
    (1, "NSW", BUSHFIRE, NATURAL, datetime.date(2019, 11, 8), None, True,
     box(150.0, -33.0, 150.2, -32.8)),
    (2, "NSW", PRESCRIBED, UNDETERMINED, datetime.date(2020, 1, 4), None, True,
     box(150.5, -33.0, 150.55, -32.95)),
    (3, "QLD", BUSHFIRE, None, datetime.date(2019, 8, 1), None, True,
     box(148.0, -26.0, 148.1, -25.9)),
    # 2020-21: July opens a season, March closes it.
    (4, "VIC", UNKNOWN, UNDETERMINED, datetime.date(2020, 7, 1), None, True,
     box(145.0, -37.0, 145.3, -36.7)),
    (5, "NSW", BUSHFIRE, NATURAL, datetime.date(2021, 3, 1), None, True,
     box(149.0, -35.0, 149.01, -34.99)),
    # 2018-19, from the extinction date: the 30th of June is still the old season.
    (6, "TAS", BUSHFIRE, None, None, datetime.date(2019, 6, 30), True,
     box(146.0, -42.0, 146.1, -41.9)),
    # Before the default first season.
    (7, "WA", BUSHFIRE, None, datetime.date(1965, 12, 1), None, True,
     box(116.0, -32.0, 116.1, -31.9)),
    # A date that cannot be believed.
    (8, "WA", BUSHFIRE, None, datetime.date(2019, 10, 1), datetime.date(2019, 9, 1),
     False, box(117.0, -32.0, 117.1, -31.9)),
]


def hectares(geometry) -> float:
    """Geodesic area of a shapely polygon in hectares, computed by PROJ."""
    area, _ = GEOD.geometry_area_perimeter(geometry)
    return abs(area) / app.SQUARE_METRES_PER_HECTARE


def reported(object_id: int) -> float:
    """The fixture's published area: the measured one, rounded to whole hectares."""
    return float(round(hectares(FIRES[object_id - 1][7])))


@pytest.fixture
def provider(db_session):
    provider = DataProvider(name=australia_csiro.PROVIDER_NAME,
                            product=australia_csiro.PROVIDER_PRODUCT,
                            full_name=australia_csiro.PROVIDER_FULL_NAME,
                            url=australia_csiro.PROVIDER_URL)
    db_session.add(provider)
    db_session.commit()
    return provider


@pytest.fixture
def fires(db_session, provider):
    """The eight fires above, stored the way the importer stores them."""
    for object_id, state, fire_type, cause, ignition, extinguish, plausible, geometry \
            in FIRES:
        start = ignition or extinguish
        db_session.add(CsiroWildfire(
            data_provider=provider,
            source_layer=australia_csiro.LAYER_NATIONAL,
            object_id=object_id,
            year=start.year,
            fire_id=str(object_id),
            fire_id_is_sentinel=False,
            state_published=state,
            state_code=state,
            fire_type=fire_type,
            cause=cause,
            ignition_date=ignition,
            extinguish_date=extinguish,
            date_source=(australia_csiro.SOURCE_IGNITION if ignition
                         else australia_csiro.SOURCE_EXTINGUISH),
            date_time_precision=australia_csiro.PRECISION_DAY,
            dates_plausible=plausible,
            area_ha=reported(object_id),
            start_date_time=datetime.datetime(start.year, start.month, start.day,
                                              tzinfo=UTC),
            time_zone=australia_csiro.DEFAULT_TIME_ZONE,
            perimeter=f"SRID=4326;{MultiPolygon([geometry]).wkt}",
        ))
    db_session.commit()
    return db_session


def run(db_session, **kwargs) -> dict:
    """Compute the reports over the fixture, keyed by state (``None`` = national)."""
    return {report.state: report for report in app.compute(db_session, logger, **kwargs)}


def season(report, start: int):
    return next(row for row in report.rows if row.season == start)


# --------------------------------------------------------------------------
# Seasons
# --------------------------------------------------------------------------

@pytest.mark.parametrize("start, label", [(2019, "2019-20"), (1999, "1999-00"),
                                          (1970, "1970-71")])
def test_a_season_is_named_by_both_years(start, label):
    assert app.season_label(start) == label


@pytest.mark.parametrize("text", ["2019", "2019-20", "2019-2020", "2019/20", " 2019 "])
def test_a_season_is_accepted_in_either_spelling(text):
    assert app.season_argument(text) == 2019


@pytest.mark.parametrize("text", ["2019-21", "19", "summer", "2019-2021"])
def test_a_season_that_is_not_one_is_refused(text):
    with pytest.raises(argparse.ArgumentTypeError):
        app.season_argument(text)


def test_a_season_straddles_new_year(fires):
    """8 November 2019 and 4 January 2020 are the same summer."""
    national = run(fires)[None]
    assert season(national, 2019).total.fires == 3


def test_july_opens_a_season_and_june_closes_one(fires):
    """1 July 2020 is 2020-21; 30 June 2019 (an extinction date) is 2018-19."""
    reports = run(fires)
    assert season(reports["VIC"], 2020).total.fires == 1
    assert season(reports["TAS"], 2018).total.fires == 1
    assert season(reports["TAS"], 2019).total.fires == 0


# --------------------------------------------------------------------------
# The figures
# --------------------------------------------------------------------------

def test_measured_is_the_geodesic_area(fires):
    nsw = run(fires)["NSW"]
    assert season(nsw, 2019).total.measured == pytest.approx(
        hectares(FIRES[0][7]) + hectares(FIRES[1][7]), rel=1e-6)


def test_reported_is_the_published_area(fires):
    nsw = run(fires)["NSW"]
    assert season(nsw, 2019).total.reported == reported(1) + reported(2)


def test_the_types_add_up_to_every_fire(fires):
    for report in run(fires).values():
        for row in report.rows:
            parts = [row.by_type.get(fire_type, app.Triplet())
                     for fire_type in app.TYPE_LABELS]
            assert sum(part.fires for part in parts) == row.total.fires
            assert sum(part.measured for part in parts) == pytest.approx(
                row.total.measured)


def test_the_nation_is_the_sum_of_the_states(fires):
    reports = run(fires)
    national = reports.pop(None)
    for row in national.rows:
        states = [next(r for r in report.rows if r.season == row.season)
                  for report in reports.values()]
        assert sum(r.total.fires for r in states) == row.total.fires
        assert sum(r.total.reported for r in states) == pytest.approx(row.total.reported)


def test_natural_cuts_across_the_types(fires):
    nsw = season(run(fires)["NSW"], 2019)
    assert nsw.natural.fires == 1
    assert nsw.natural.measured == pytest.approx(hectares(FIRES[0][7]), rel=1e-6)


def test_natural_is_empty_where_no_cause_is_published(fires):
    """QLD publishes no cause: an empty cell, not a zero that claims an answer."""
    reports = run(fires)
    assert season(reports["QLD"], 2019).natural is None
    assert season(reports["QLD"], 2019).values[-4:-1] == ("", "", "")
    assert reports["QLD"].rows[-1].natural is None


def test_natural_is_zero_where_causes_are_published_and_none_is_natural(fires):
    """VIC 2020-21 publishes an undetermined cause: zero natural fires is an answer."""
    vic = season(run(fires)["VIC"], 2020)
    assert vic.natural == app.Triplet()
    assert vic.values[-4:-1] == ("0", "0.00", "0.00")


def test_fires_with_cause_counts_the_readable_causes(fires):
    """2019-20: NSW's natural and undetermined fires have one, QLD's has none."""
    reports = run(fires)
    assert season(reports[None], 2019).with_cause == 2
    assert season(reports["NSW"], 2019).with_cause == 2
    assert season(reports["QLD"], 2019).with_cause == 0
    assert season(reports[None], 2019).values[-1] == "2"


def test_fires_with_cause_is_summed_in_the_total(fires):
    reports = run(fires, first=app.DEFAULT_FIRST_SEASON)
    assert reports[None].rows[-1].with_cause == 4, "fires 1, 2, 4 and 5"
    assert reports["TAS"].rows[-1].with_cause == 0


# --------------------------------------------------------------------------
# Rows and scope
# --------------------------------------------------------------------------

def test_the_default_starts_in_1970_and_leaves_out_bad_dates(fires):
    national = run(fires, first=app.DEFAULT_FIRST_SEASON)[None]
    assert national.rows[-2].season == 1970
    assert national.rows[-1].total.fires == 6, "the 1965 fire and the bad date left out"


def test_every_season_can_be_asked_for(fires):
    national = run(fires)[None]
    assert national.rows[-2].season == 1965
    assert national.rows[-1].total.fires == 7


def test_bad_dates_can_be_counted_on_request(fires):
    national = run(fires, include_implausible=True)[None]
    assert season(national, 2019).total.fires == 4


def test_seasons_without_fires_are_zero_rows_and_every_file_lines_up(fires):
    reports = run(fires, first=app.DEFAULT_FIRST_SEASON)
    expected = list(range(2020, 1969, -1))
    for report in reports.values():
        assert [row.season for row in report.rows[:-1]] == expected
    assert season(reports[None], 1990).total == app.Triplet()


def test_the_seasons_are_newest_first_and_the_total_last(fires):
    rows = run(fires)[None].rows
    assert rows[0].season == 2020
    assert rows[-1].is_total and rows[-1].season_label == "Total"


def test_a_state_with_no_fires_gets_no_report(fires):
    assert set(run(fires, first=app.DEFAULT_FIRST_SEASON)) == {
        None, "NSW", "QLD", "VIC", "TAS"}, "WA's only fires are out of scope"


def test_one_season(fires):
    reports = run(fires, first=2019, last=2019)
    assert [row.season for row in reports[None].rows] == [2019, None]
    assert set(reports) == {None, "NSW", "QLD"}


def test_nothing_matched_is_no_report(fires):
    assert run(fires, first=2030) == {}


# --------------------------------------------------------------------------
# The command line
# --------------------------------------------------------------------------

def test_an_output_is_required():
    with pytest.raises(SystemExit):
        app.parse_arguments([])


def test_the_defaults():
    args = app.parse_arguments(["--csv", "a.csv"])
    assert args.from_season == app.DEFAULT_FIRST_SEASON == 1970
    assert args.to_season is None
    assert not args.include_implausible


def test_one_season_sets_both_ends():
    args = app.parse_arguments(["--csv", "a.csv", "--season", "2019-20"])
    assert (args.from_season, args.to_season) == (2019, 2019)


def test_all_seasons_removes_the_default_start():
    assert app.parse_arguments(["--csv", "a.csv", "--all-seasons"]).from_season is None


def test_a_backwards_range_is_refused():
    with pytest.raises(SystemExit):
        app.parse_arguments(["--csv", "a.csv", "--from-season", "2020",
                             "--to-season", "2019"])


# --------------------------------------------------------------------------
# The files
# --------------------------------------------------------------------------

def test_a_state_file_sits_beside_the_national_one(tmp_path):
    assert app.output_path(tmp_path / "au.csv", None) == tmp_path / "au.csv"
    assert app.output_path(tmp_path / "au.csv", "NSW") == tmp_path / "au_NSW.csv"


def test_the_csv_has_a_season_and_five_triplets(tmp_path, fires):
    national = run(fires, first=2019, last=2019)[None]
    app.write_csv(national, tmp_path / "au.csv", logger)
    rows = list(csv.reader((tmp_path / "au.csv").open()))
    assert rows[0] == list(app.COLUMNS) and len(rows[0]) == 17
    assert rows[0][-1] == app.WITH_CAUSE_LABEL
    assert rows[0][:4] == ["Season", "Fires", "Reported (ha)", "Measured (ha)"]
    assert rows[1][0] == "2019-20" and rows[1][1] == "3"
    assert rows[-1][0] == "Total"


def test_the_docx_is_landscape_with_a_grouped_heading(tmp_path, fires):
    from docx import Document

    national = run(fires, first=2019, last=2019)[None]
    app.write_docx(national, tmp_path / "au.docx", logger)
    document = Document(str(tmp_path / "au.docx"))
    section = document.sections[0]
    assert section.page_width > section.page_height
    table = document.tables[0]
    assert len(table.columns) == 17
    assert table.cell(0, 16).text == table.cell(1, 16).text == app.WITH_CAUSE_LABEL
    assert table.cell(0, 1).text == app.ALL_LABEL
    assert table.cell(0, 13).text == app.NATURAL_LABEL
    assert table.cell(1, 2).text == "Reported (ha)"
    assert table.cell(2, 0).text == "2019-20"


def test_the_run_writes_a_file_per_state(tmp_path, fires, monkeypatch):
    args = app.parse_arguments(["--csv", str(tmp_path / "au.csv"),
                                "--docx", str(tmp_path / "au.docx")])
    monkeypatch.setattr(app, "Session", lambda engine: _Borrowed(fires))
    app.report(args, None, logger)
    written = sorted(path.name for path in tmp_path.iterdir())
    assert written == sorted(f"au{suffix}.{kind}"
                             for suffix in ("", "_NSW", "_QLD", "_VIC", "_TAS")
                             for kind in ("csv", "docx"))


class _Borrowed:
    """The test session, lent to :func:`app.report` without being closed by it."""

    def __init__(self, session):
        self.session = session

    def __enter__(self):
        return self.session

    def __exit__(self, *exc):
        return False
