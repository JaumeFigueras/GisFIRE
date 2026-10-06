#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Wildfire counts by cause for the New Zealand Wildfire Occurrence Database.

Reports, per country and year, how many fires the WOD records, how many have a
**determined** cause, how they split between lightning, other natural causes and
people, and how much of the recorded area the lightning fires account for::

    Country      Year   Fires  Determined  Lightning  Lightning (%)  Natural  Human  Unknown  Lightning (ha)  Lightning (% of ha)
    New Zealand  2024    4703         ...        ...            ...      ...    ...      ...            ...                 ...
    New Zealand  Total    ...

::

    python3 -m src.apps.statistics.wildfires.new_zealand_wod.wildfire_causes --csv causes.csv
    python3 -m src.apps.statistics.wildfires.new_zealand_wod.wildfire_causes \\
        --financial-year --docx causes.docx

At least one of ``--csv`` and ``--docx`` is required. The application only reads.

The companion of :mod:`~src.apps.statistics.wildfires.new_zealand_wod.
wildfire_statistics`, over the same fires and the same scope, so the ``Country``,
``Year`` and ``Fires`` columns of the two reports agree row for row. The years query,
the scope, the country resolution and the country order are imported from it rather
than copied.

This dataset does have lightning
---------------------------------

Unlike the Canadian or Portuguese archives, whose *natural* class is only a proxy, the
WOD's causes are specific: 538 records name lightning outright (FENZ ``Lightning
discharge``, DOC ``Lightning`` and ``Lightening Strike``). They are the ``lightning``
class of :data:`~src.providers.new_zealand_wod.FIRE_CAUSES`, and **natural** is kept
apart from it — geothermal or volcanic, five records — so the ``Lightning`` column
counts lightning and nothing else.

New Zealand's fires are almost entirely human-caused, so the lightning share is a
fraction of a percent of the fires. Its share of the area is the more interesting
number: a lightning fire in the back country can burn for longer before anyone reaches
it.

Why the denominator is the **determined** fires
------------------------------------------------

``unknown`` gathers records whose cause was not recorded or not classifiable — about a
sixth of the archive, and not evenly: FENZ's *Information not recorded/Unknown*
alone is around ten percent of its records. The percentages are therefore of the
fires whose cause is known — lightning, natural or human — as in the other cause
reports, and ``Unknown`` is shown as a column so that the reader can see what was
left out. Where nothing is determined the percentage cell is **empty**, not zero.

The hectares are the recorded ``AreaHa``; the perimeters are not used here.
"""

from __future__ import annotations

import argparse
import csv
import logging
import math
import sys

from dataclasses import dataclass
from pathlib import Path

from sqlalchemy import Engine
from sqlalchemy import Select
from sqlalchemy import func
from sqlalchemy import select
from sqlalchemy.orm import Session

from src.apps.imports import common
from src.apps.statistics.wildfires.new_zealand_wod.wildfire_statistics import COUNTRY_NAME
from src.apps.statistics.wildfires.new_zealand_wod.wildfire_statistics import (
    COUNTRY_SOURCE_FILED,
)
from src.apps.statistics.wildfires.new_zealand_wod.wildfire_statistics import (
    COUNTRY_SOURCE_GEOMETRY,
)
from src.apps.statistics.wildfires.new_zealand_wod.wildfire_statistics import WOD
from src.apps.statistics.wildfires.new_zealand_wod.wildfire_statistics import (
    add_common_arguments,
)
from src.apps.statistics.wildfires.new_zealand_wod.wildfire_statistics import (
    country_columns,
)
from src.apps.statistics.wildfires.new_zealand_wod.wildfire_statistics import (
    coverage_paragraph,
)
from src.apps.statistics.wildfires.new_zealand_wod.wildfire_statistics import (
    ordered_countries,
)
from src.apps.statistics.wildfires.new_zealand_wod.wildfire_statistics import (
    outside_count,
)
from src.apps.statistics.wildfires.new_zealand_wod.wildfire_statistics import period_column
from src.apps.statistics.wildfires.new_zealand_wod.wildfire_statistics import period_label
from src.apps.statistics.wildfires.new_zealand_wod.wildfire_statistics import (
    run_application,
)
from src.apps.statistics.wildfires.new_zealand_wod.wildfire_statistics import (
    scope_conditions,
)
from src.apps.statistics.wildfires.new_zealand_wod.wildfire_statistics import summarise
from src.apps.statistics.wildfires.new_zealand_wod.wildfire_statistics import write_table
from src.apps.statistics.wildfires.new_zealand_wod.wildfire_statistics import years_query
from src.providers import new_zealand_wod

#: The report's columns, in order. The first three are the companion report's.
COLUMNS = ("Country", "Year", "Fires", "Determined", "Lightning", "Lightning (%)",
           "Natural", "Human", "Unknown", "Lightning (ha)", "Lightning (% of ha)")

#: The classes whose cause somebody determined: the denominator.
DETERMINED_CAUSES = (new_zealand_wod.CAUSE_LIGHTNING, new_zealand_wod.CAUSE_NATURAL,
                     new_zealand_wod.CAUSE_HUMAN)


def counts_query(year: int, financial_year: bool = False,
                 country_source: str = COUNTRY_SOURCE_FILED,
                 source: str | None = None) -> Select:
    """Build the counting query for one year.

    Returns
    -------
    Select
        ``country, fires, lightning, natural, human, unknown, lightning_hectares,
        determined_hectares``: one row per country.

    Notes
    -----
    Every count in one pass. The cause is what is counted, so only the ``source`` half
    of the companion's scope applies — which is what keeps ``Fires`` equal to that
    report's under the same options. A recorded area that is ``NULL`` adds nothing to
    the hectares and the fire still counts.
    """
    country_name, joins = country_columns(country_source)
    cause = WOD.c.fire_cause
    area = WOD.c.area_ha
    determined = cause.in_(DETERMINED_CAUSES)
    lightning = cause == new_zealand_wod.CAUSE_LIGHTNING

    counts = (
        select(
            country_name.label("country"),
            func.count().label("fires"),
            func.count().filter(lightning).label("lightning"),
            func.count().filter(cause == new_zealand_wod.CAUSE_NATURAL).label("natural"),
            func.count().filter(cause == new_zealand_wod.CAUSE_HUMAN).label("human"),
            func.count().filter(cause == new_zealand_wod.CAUSE_UNKNOWN).label("unknown"),
            func.coalesce(func.sum(area).filter(lightning), 0.0).label("lightning_hectares"),
            func.coalesce(func.sum(area).filter(determined),
                          0.0).label("determined_hectares"),
        )
        .select_from(WOD)
        .where(period_column(financial_year) == year)
    )
    for target, condition in joins:
        counts = counts.join(target, condition)
    for condition in scope_conditions(None, source):
        counts = counts.where(condition)
    return counts.group_by(country_name)


@dataclass(frozen=True)
class Row:
    """One line of the report.

    Attributes
    ----------
    country : str
        The country the fire counts towards.
    year : int or None
        The year, or the year the fire year starts in, or ``None`` for a summary row.
    fires, lightning, natural, human, unknown : int
        Every fire, and each cause class. They add up: ``fires`` is the sum of the four.
    lightning_hectares, determined_hectares : float
        The recorded area of the lightning fires, and of all the determined ones.
    financial_year : bool
        Whether :attr:`year` is a fire year.
    """

    country: str
    year: int | None
    fires: int
    lightning: int
    natural: int
    human: int
    unknown: int
    lightning_hectares: float
    determined_hectares: float
    financial_year: bool = False

    @property
    def is_total(self) -> bool:
        return self.year is None

    @property
    def year_label(self) -> str:
        return period_label(self.year, self.financial_year)

    @property
    def determined(self) -> int:
        """The fires with a known cause: the denominator."""
        return self.lightning + self.natural + self.human

    @property
    def share(self) -> float | None:
        """Lightning as a percentage of the determined fires, or ``None`` if none is."""
        return 100.0 * self.lightning / self.determined if self.determined else None

    @property
    def area_share(self) -> float | None:
        """Lightning's share of the determined fires' recorded area."""
        if not self.determined_hectares:
            return None
        return 100.0 * self.lightning_hectares / self.determined_hectares

    @staticmethod
    def _percentage(value: float | None) -> str:
        return "" if value is None else f"{value:.2f}"

    @property
    def values(self) -> tuple[str, ...]:
        """The row as the CSV writes it, in :data:`COLUMNS` order."""
        return (self.country, self.year_label, str(self.fires), str(self.determined),
                str(self.lightning), self._percentage(self.share), str(self.natural),
                str(self.human), str(self.unknown), f"{self.lightning_hectares:.4f}",
                self._percentage(self.area_share))

    @property
    def readable_values(self) -> tuple[str, ...]:
        """The row as the Word document writes it: numbers with separators."""
        return (self.country, self.year_label, f"{self.fires:,}", f"{self.determined:,}",
                f"{self.lightning:,}", self._percentage(self.share), f"{self.natural:,}",
                f"{self.human:,}", f"{self.unknown:,}", f"{self.lightning_hectares:,.2f}",
                self._percentage(self.area_share))


def combine(rows: list[Row], country: str, year: int | None = None) -> Row:
    """One row summarising several: counts and hectares added, shares recomputed."""
    return Row(
        country=country, year=year,
        fires=sum(row.fires for row in rows),
        lightning=sum(row.lightning for row in rows),
        natural=sum(row.natural for row in rows),
        human=sum(row.human for row in rows),
        unknown=sum(row.unknown for row in rows),
        lightning_hectares=math.fsum(row.lightning_hectares for row in rows),
        determined_hectares=math.fsum(row.determined_hectares for row in rows),
        financial_year=rows[0].financial_year,
    )


def parse_arguments(argv: list[str] | None = None) -> argparse.Namespace:
    """Parse the command line."""
    parser = argparse.ArgumentParser(
        description="Wildfire counts by cause for the New Zealand Wildfire Occurrence "
                    "Database, with the lightning fires' share of the recorded area.",
        epilog="The percentages are of the fires whose cause is known. Database settings "
               "not given here are read from the environment (.env).",
    )
    add_common_arguments(parser)
    arguments = parser.parse_args(argv)
    if arguments.csv is None and arguments.docx is None:
        parser.error("nothing to write: pass --csv, --docx, or both")
    return arguments


def compute(session: Session, year: int | None, logger: logging.Logger,
            financial_year: bool = False, country_source: str = COUNTRY_SOURCE_FILED,
            source: str | None = None) -> list[Row]:
    """Count the fires a year at a time, returning the report's rows in order."""
    if year is not None:
        years = [year]
    else:
        years = list(session.scalars(years_query(financial_year, None, source)))

    measured: list[Row] = []
    for index, counting in enumerate(years, start=1):
        label = period_label(counting, financial_year)
        with common.Spinner(f"Counting the WOD fires by cause ({label}: {index} of "
                            f"{len(years)})", logger):
            measured += [
                Row(country=record.country, year=counting, fires=record.fires,
                    lightning=record.lightning, natural=record.natural,
                    human=record.human, unknown=record.unknown,
                    lightning_hectares=float(record.lightning_hectares),
                    determined_hectares=float(record.determined_hectares),
                    financial_year=financial_year)
                for record in session.execute(
                    counts_query(counting, financial_year, country_source, source))
            ]

    if country_source == COUNTRY_SOURCE_GEOMETRY:
        outside = session.scalar(outside_count(year, financial_year, source=source)) or 0
        if outside:
            logger.warning("%d fire(s) have a point inside no country and are not "
                           "counted; --country-source filed counts them", outside)

    countries = ordered_countries(session, {row.country for row in measured})
    rows = summarise(measured, countries, combine)
    if measured:
        total = combine(measured, COUNTRY_NAME)
        logger.info("%d of %d fire(s) have a known cause; %d are lightning (%s%%), "
                    "accounting for %s%% of the determined fires' recorded area",
                    total.determined, total.fires, total.lightning,
                    Row._percentage(total.share), Row._percentage(total.area_share))
    return rows


def write_csv(rows: list[Row], path: Path, logger: logging.Logger) -> None:
    """Write the report as CSV; an empty percentage is no answer, not zero."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(COLUMNS)
        for row in rows:
            writer.writerow(row.values)
    logger.info("Wrote %s", path)


def write_docx(rows: list[Row], path: Path, logger: logging.Logger, year: int | None,
               financial_year: bool, country_source: str, source: str | None) -> None:
    """Write the report as a Word document, its caveats first."""
    from docx import Document

    document = Document()
    document.add_heading("WOD wildfires by cause (New Zealand)", level=1)
    scope = [f"year: {period_label(year, financial_year)}" if year is not None
             else "all years"]
    scope.append("fire years (1 July to 30 June)" if financial_year else "calendar years")
    if source is not None:
        scope.append(f"only {source} records")
    if country_source == COUNTRY_SOURCE_GEOMETRY:
        scope.append("only fires whose point falls inside a country")
    document.add_paragraph(
        "Counts of the WOD records by cause class. Lightning is lightning as recorded "
        "(FENZ 'Lightning discharge', DOC 'Lightning' and 'Lightening Strike'); Natural "
        "is any other natural cause, geothermal or volcanic. The percentages are of the "
        "fires whose cause is known; Unknown is shown so that what was left out is "
        "visible. Hectares are the recorded AreaHa. Scope: " + "; ".join(scope) + ".")
    document.add_paragraph(coverage_paragraph())
    write_table(document, COLUMNS, rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    logger.info("Wrote %s", path)


def report(args: argparse.Namespace, engine: Engine, logger: logging.Logger) -> list[Row]:
    """Count the causes and write whichever outputs were asked for."""
    with Session(engine) as session:
        rows = compute(session, args.year, logger, args.financial_year,
                       args.country_source, args.source)
    if not rows:
        raise RuntimeError(
            "No wildfires matched. Check --year and --source, and that the WOD is "
            f"imported — the import reads from {new_zealand_wod.FIRST_YEAR} on unless "
            "run with --all-years.")
    if args.csv is not None:
        write_csv(rows, args.csv, logger)
    if args.docx is not None:
        write_docx(rows, args.docx, logger, args.year, args.financial_year,
                   args.country_source, args.source)
    return rows


def main(argv: list[str] | None = None) -> int:
    return run_application(argv, parse_arguments, report, "wod-causes")


if __name__ == "__main__":  # pragma nocover
    sys.exit(main())
