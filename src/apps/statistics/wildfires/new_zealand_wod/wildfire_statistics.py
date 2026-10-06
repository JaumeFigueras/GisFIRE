#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Burnt-area statistics for the New Zealand Wildfire Occurrence Database.

Reports, per country and year, how many fires the WOD records and the smallest,
largest and total area burnt, in hectares — and, this dataset's own columns, how many
of those fires each agency recorded and how many have a perimeter::

    Country      Year   Fires  Minimum (ha)  Maximum (ha)  Total (ha)  FENZ  DOC  Perimeters
    New Zealand  2024    4703          0.00       ...         ...      4703    0        ...
    ...
    New Zealand  Total  ...

Run it over everything, or narrow it to one year, to the fires above a size, to one
cause or to one agency; or count by fire year rather than calendar year::

    python3 -m src.apps.statistics.wildfires.new_zealand_wod.wildfire_statistics \\
        --csv burnt.csv
    python3 -m src.apps.statistics.wildfires.new_zealand_wod.wildfire_statistics \\
        --financial-year --docx fire-years.docx
    python3 -m src.apps.statistics.wildfires.new_zealand_wod.wildfire_statistics \\
        --cause lightning --csv lightning.csv
    python3 -m src.apps.statistics.wildfires.new_zealand_wod.wildfire_statistics \\
        --area-method perimeter --csv mapped.csv

At least one of ``--csv`` and ``--docx`` is required. The application only reads.
Database settings come from the environment (``.env``, see :mod:`src.settings`); every
one of them can be overridden with a command-line argument.

Its first six columns are the other burnt-area reports', in their order, so the CSVs
can be concatenated on them; the last three are this dataset's own.

Which area
----------

``--area-method`` chooses:

``reported`` (the default)
    :attr:`~src.providers.new_zealand_wod.wildfire.WodWildfire.area_ha`, the area the
    agency recorded. Every fire has one, a 0 is a real answer, and it is the only area
    most of the archive has.
``perimeter``
    The area of the published perimeter, measured on NZTM2000 — so **only the 2,089
    fires with one are counted**. The paper notes that FENZ perimeters are often
    traced from helicopters and include unburnt islands, while DOC's were mostly
    digitised from paper maps.

The two are not interchangeable and the report never mixes them: a row is either
every fire's reported area or the perimeter fires' measured one.

.. warning::

   **Most of these fires are tiny.** The median FENZ record is 0.0005 ha — five square
   metres — because every vegetation fire FENZ attended is a record. A ``Fires`` count
   here is a count of incidents, most of them a few square metres of grass, and is not
   comparable with a satellite or a perimeter archive's count. ``--min-area`` is how to
   ask about fires of a size.

Which years are comparable
--------------------------

* **Before 1998** the record is DOC's alone and, in the paper's words, does not give
  good overall coverage; the import leaves those years out unless asked for them.
* **1998-2011** both agencies, deduplicated by the compilers (same day, within 600 m:
  the DOC record kept).
* **From 2012** FENZ alone. The ``FENZ`` and ``DOC`` columns are this caveat made
  visible.
* There are **no perimeters for 2012-2016**, so a ``--area-method perimeter`` series
  has a hole there that is not an absence of fire.

Calendar year or fire year
---------------------------

By default a fire counts towards its published calendar ``Year``. ``--financial-year``
counts it towards its New Zealand fire year instead — 1 July to 30 June, labelled
``1998/1999`` — which is how FENZ and the annual summaries count, and which does not
split a southern summer in two. ``--year`` then names the year the fire year starts in.

Which country a fire counts towards
------------------------------------

``--country-source`` chooses, as in the NFDB report, but the default here is the
other way round:

``filed`` (the default)
    Every WOD record is a New Zealand agency's record of a New Zealand fire, so the
    ``Country`` column is the constant :data:`COUNTRY_NAME`. The 92 points in the sea
    are such fires, recorded at a wrong location — the paper keeps them, and so does
    this report.
``geometry``
    The published point tested against the OCHA country polygons. The points in the
    sea are left out, and the log says how many.

One year at a time
------------------

One statement per year, the shape every burnt-area report in the project has, so
that ``geometry``'s point-in-polygon tests release their memory as each year ends.
The summary rows are computed from the years' rows, which decompose exactly.
"""

from __future__ import annotations

import argparse
import csv
import logging
import math
import os
import sys

from dataclasses import dataclass
from pathlib import Path

from sqlalchemy import ColumnElement
from sqlalchemy import Engine
from sqlalchemy import Select
from sqlalchemy import create_engine
from sqlalchemy import func
from sqlalchemy import literal
from sqlalchemy import select
from sqlalchemy import true
from sqlalchemy.orm import Session

import src.settings  # noqa: F401  (imported for the side effect of loading .env)

from src.apps.imports import common
from src.data_model.geography.admin_boundary import AdminBoundary
from src.data_model.ignition import Ignition
from src.providers import new_zealand_wod
from src.providers.new_zealand_wod.wildfire import WodWildfire
from src.providers.ocha.admin_boundary import OchaAdminBoundary

#: Label used in the ``Year`` column for a summary row.
TOTAL_LABEL = "Total"

#: The six columns this report shares with the other burnt-area reports, in order.
SHARED_COLUMNS = ("Country", "Year", "Fires", "Minimum (ha)", "Maximum (ha)", "Total (ha)")

#: The report's columns: the shared six, then how many of the fires each agency
#: recorded and how many have a perimeter.
COLUMNS = SHARED_COLUMNS + ("FENZ", "DOC", "Perimeters")

#: Index of the first column that holds a number, right-aligned in the Word table.
FIRST_NUMERIC_COLUMN = 2

#: The country every WOD record is filed in: the ``Country`` column under
#: ``--country-source filed``. Spelled as the OCHA boundaries spell it.
COUNTRY_NAME = "New Zealand"

#: The two ways of deciding which country a fire counts towards. ``filed`` is the
#: default here; see the module docstring.
COUNTRY_SOURCE_FILED = "filed"
COUNTRY_SOURCE_GEOMETRY = "geometry"
COUNTRY_SOURCES = (COUNTRY_SOURCE_FILED, COUNTRY_SOURCE_GEOMETRY)

#: Administrative level of a country in ``admin_boundary``.
COUNTRY_LEVEL = 0

#: The two areas a fire can be measured by. See the module docstring.
AREA_REPORTED = "reported"
AREA_PERIMETER = "perimeter"
AREA_METHODS = (AREA_REPORTED, AREA_PERIMETER)

#: The four cause classes, as ``--cause`` accepts them: the provider's own values.
CAUSES = new_zealand_wod.FIRE_CAUSES

#: The two agencies, as ``--source`` accepts them.
SOURCES = new_zealand_wod.SOURCES

#: Square metres in a hectare.
SQUARE_METRES_PER_HECTARE = 10_000.0

WOD = WodWildfire.__table__


def period_column(financial_year: bool) -> ColumnElement:
    """The column a fire's year is read from: calendar ``year`` or ``financial_year``."""
    return WOD.c.financial_year if financial_year else WOD.c.year


def period_label(year: int | None, financial_year: bool) -> str:
    """How a year is written in the ``Year`` column: ``2005``, ``2005/2006`` or ``Total``."""
    if year is None:
        return TOTAL_LABEL
    return f"{year}/{year + 1}" if financial_year else str(year)


def area_column(area_method: str) -> ColumnElement:
    """A fire's area in hectares under ``area_method``.

    Raises
    ------
    ValueError
        If ``area_method`` is not one of :data:`AREA_METHODS`.

    Notes
    -----
    Under ``perimeter`` this is ``NULL`` for a fire without one, and every query filters
    ``NULL`` areas out, so the point fires drop out of that report rather than counting
    as zero.
    """
    if area_method == AREA_REPORTED:
        return WOD.c.area_ha
    if area_method == AREA_PERIMETER:
        return func.ST_Area(WOD.c.perimeter_nztm) / SQUARE_METRES_PER_HECTARE
    raise ValueError(f"unknown area method {area_method!r}; expected one of "
                     f"{', '.join(AREA_METHODS)}")


def country_columns(source: str) -> tuple[ColumnElement, list]:
    """Where a fire's country comes from, as ``(name, joins)``.

    ``filed`` is the constant :data:`COUNTRY_NAME` and needs no join. ``geometry``
    joins the fire's point — every WOD fire has one — and tests it against the OCHA
    countries, ``LATERAL ... LIMIT 1`` so a point on a shared border cannot become two
    rows. A point in no country drops out of the inner join; :func:`outside_count`
    counts those.

    Raises
    ------
    ValueError
        If ``source`` is not one of :data:`COUNTRY_SOURCES`.
    """
    if source == COUNTRY_SOURCE_FILED:
        return literal(COUNTRY_NAME), []
    if source != COUNTRY_SOURCE_GEOMETRY:
        raise ValueError(f"unknown country source {source!r}; expected one of "
                         f"{', '.join(COUNTRY_SOURCES)}")
    ignition = Ignition.__table__
    ocha_boundary = OchaAdminBoundary.__table__
    containing = (
        select(AdminBoundary.name.label("name"))
        .select_from(AdminBoundary)
        .join(ocha_boundary, ocha_boundary.c.id == AdminBoundary.id)
        .where(AdminBoundary.level == COUNTRY_LEVEL)
        .where(func.ST_Contains(AdminBoundary.geometry, ignition.c.geometry))
        .limit(1)
        .lateral("containing")
    )
    return containing.c.name, [(ignition, ignition.c.id == WOD.c.ignition_id),
                               (containing, true())]


def scope_conditions(cause: str | None = None,
                     source: str | None = None) -> list[ColumnElement]:
    """Every condition narrowing which fires the report is of.

    Built once and applied to the years query, each year's statement and the audit, so
    the three cannot drift apart. Shared with the causes report, which applies only the
    ``source`` half: the cause is what that report counts.

    Raises
    ------
    ValueError
        For a cause or a source outside the provider's vocabulary.
    """
    conditions = []
    if cause is not None:
        if cause not in CAUSES:
            raise ValueError(f"unknown cause {cause!r}; expected one of {', '.join(CAUSES)}")
        conditions.append(WOD.c.fire_cause == cause)
    if source is not None:
        if source not in SOURCES:
            raise ValueError(f"unknown source {source!r}; expected one of "
                             f"{', '.join(SOURCES)}")
        conditions.append(WOD.c.source == source)
    return conditions


def years_query(financial_year: bool = False, cause: str | None = None,
                source: str | None = None) -> Select:
    """The years the stored fires fall in, newest first, under the scope.

    ``DISTINCT`` and not a range, so a gap in the record is a gap in the report.
    """
    period = period_column(financial_year)
    query = select(period.label("year")).distinct().order_by(period.desc())
    for condition in scope_conditions(cause, source):
        query = query.where(condition)
    return query


def statistics_query(year: int, financial_year: bool = False,
                     area_method: str = AREA_REPORTED,
                     min_area: float | None = None,
                     country_source: str = COUNTRY_SOURCE_FILED,
                     cause: str | None = None,
                     source: str | None = None) -> Select:
    """Build the statistics query for one year.

    Returns
    -------
    Select
        A query yielding ``country, minimum, maximum, total, fires, fenz, doc,
        perimeters``: one row per country with a fire in ``year``.

    Notes
    -----
    Selected from ``wod_wildfire`` by table, so no polymorphic join is added: the
    year, the area, the cause, the agency and the perimeter all live there. A fire
    with no area under the chosen method — no perimeter, under ``perimeter`` — is
    filtered out before the aggregates. ``min_area`` likewise selects fires, not years.
    """
    country_name, joins = country_columns(country_source)
    area = area_column(area_method)

    fires = (
        select(country_name.label("country"),
               area.label("hectares"),
               WOD.c.source.label("source"),
               WOD.c.perimeter_nztm.is_not(None).label("has_perimeter"))
        .select_from(WOD)
        .where(period_column(financial_year) == year)
        .where(area.is_not(None))
    )
    for target, condition in joins:
        fires = fires.join(target, condition)
    for condition in scope_conditions(cause, source):
        fires = fires.where(condition)

    fire = fires.subquery("fire")
    statistics = (
        select(
            fire.c.country,
            func.min(fire.c.hectares).label("minimum"),
            func.max(fire.c.hectares).label("maximum"),
            func.sum(fire.c.hectares).label("total"),
            func.count().label("fires"),
            func.count().filter(fire.c.source == new_zealand_wod.SOURCE_FENZ).label("fenz"),
            func.count().filter(fire.c.source == new_zealand_wod.SOURCE_DOC).label("doc"),
            func.count().filter(fire.c.has_perimeter).label("perimeters"),
        )
        .group_by(fire.c.country)
    )
    if min_area is not None:
        statistics = statistics.where(fire.c.hectares >= min_area)
    return statistics


def outside_count(year: int | None, financial_year: bool = False,
                  area_method: str = AREA_REPORTED, min_area: float | None = None,
                  cause: str | None = None, source: str | None = None) -> Select:
    """How many fires in scope ``--country-source geometry`` leaves out.

    Every WOD fire has a point, so the only thing ``geometry`` can drop is a point in
    no country — the points in the sea. Counted under the report's own scope so the
    number and the ``Fires`` column account for the same fires.
    """
    ignition = Ignition.__table__
    ocha_boundary = OchaAdminBoundary.__table__
    located = (
        select(AdminBoundary.id)
        .select_from(AdminBoundary)
        .join(ocha_boundary, ocha_boundary.c.id == AdminBoundary.id)
        .where(AdminBoundary.level == COUNTRY_LEVEL)
        .where(func.ST_Contains(AdminBoundary.geometry, ignition.c.geometry))
        .exists()
    )
    area = area_column(area_method)
    query = (
        select(func.count())
        .select_from(WOD)
        .join(ignition, ignition.c.id == WOD.c.ignition_id)
        .where(area.is_not(None))
        .where(~located)
    )
    if year is not None:
        query = query.where(period_column(financial_year) == year)
    if min_area is not None:
        query = query.where(area >= min_area)
    for condition in scope_conditions(cause, source):
        query = query.where(condition)
    return query


@dataclass(frozen=True)
class Row:
    """One line of the report.

    Attributes
    ----------
    country : str
        :data:`COUNTRY_NAME`, or under ``geometry`` the country containing the point.
    year : int or None
        The calendar year, or the year the fire year starts in, or ``None`` for a
        summary row.
    minimum, maximum, total : float
        Smallest fire, largest fire and their sum, in hectares.
    fires, fenz, doc, perimeters : int
        How many fires the figures are computed from, how many of them each agency
        recorded, and how many have a perimeter.
    financial_year : bool
        Whether :attr:`year` is a fire year, which decides how it is written.
    """

    country: str
    year: int | None
    minimum: float
    maximum: float
    total: float
    fires: int
    fenz: int
    doc: int
    perimeters: int
    financial_year: bool = False

    @property
    def is_total(self) -> bool:
        return self.year is None

    @property
    def year_label(self) -> str:
        return period_label(self.year, self.financial_year)

    @property
    def values(self) -> tuple[str, ...]:
        """The row as the CSV writes it, in :data:`COLUMNS` order."""
        return (self.country, self.year_label, str(self.fires),
                f"{self.minimum:.4f}", f"{self.maximum:.4f}", f"{self.total:.4f}",
                str(self.fenz), str(self.doc), str(self.perimeters))

    @property
    def readable_values(self) -> tuple[str, ...]:
        """The row as the Word document writes it: numbers with separators."""
        return (self.country, self.year_label, f"{self.fires:,}",
                f"{self.minimum:,.4f}", f"{self.maximum:,.2f}", f"{self.total:,.2f}",
                f"{self.fenz:,}", f"{self.doc:,}", f"{self.perimeters:,}")


def combine(rows: list[Row], country: str, year: int | None = None) -> Row:
    """One row summarising several: every figure decomposes over a partition."""
    return Row(
        country=country,
        year=year,
        minimum=min(row.minimum for row in rows),
        maximum=max(row.maximum for row in rows),
        total=math.fsum(row.total for row in rows),
        fires=sum(row.fires for row in rows),
        fenz=sum(row.fenz for row in rows),
        doc=sum(row.doc for row in rows),
        perimeters=sum(row.perimeters for row in rows),
        financial_year=rows[0].financial_year,
    )


def ordered_countries(session: Session, names: set[str]) -> list[str]:
    """The countries of the report in the database's collation order.

    A name the boundaries do not hold — the constant ``New Zealand`` when the OCHA
    boundaries are not imported — is appended rather than dropped.
    """
    if not names:
        return []
    ordered = list(session.scalars(
        select(AdminBoundary.name).distinct()
        .where(AdminBoundary.level == COUNTRY_LEVEL)
        .where(AdminBoundary.name.in_(names))
        .order_by(AdminBoundary.name)
    ))
    return ordered + sorted(names - set(ordered))


def summarise(measured: list, countries: list[str], combine_rows=combine) -> list:
    """Each country, its years newest first and then its summary row."""
    report = []
    for name in countries:
        rows = [row for row in measured if row.country == name]
        if not rows:
            continue
        report += sorted(rows, key=lambda row: row.year, reverse=True)
        report.append(combine_rows(rows, name, None))
    return report


def hectares(text: str) -> float:
    """Argparse type for ``--min-area``: a finite, non-negative number of hectares."""
    try:
        area = float(text)
    except ValueError:
        raise argparse.ArgumentTypeError(f"{text!r} is not a number of hectares")
    if not math.isfinite(area):
        raise argparse.ArgumentTypeError(f"{text!r} is not a finite number of hectares")
    if area < 0:
        raise argparse.ArgumentTypeError(f"a burnt area cannot be negative, and {area:g} "
                                         f"ha is")
    return area


def add_common_arguments(parser: argparse.ArgumentParser) -> None:
    """The options this report shares with the causes report."""
    parser.add_argument("-y", "--year", type=int,
                        help="restrict to one year; with --financial-year, the year the "
                             "fire year starts in")
    parser.add_argument("--financial-year", action="store_true",
                        help="count by New Zealand fire year (1 July to 30 June, "
                             "written 2005/2006) rather than calendar year, as FENZ does")
    parser.add_argument("--source", default=None, choices=SOURCES,
                        help="restrict to one agency's records. DOC's stop after 2011 "
                             "and FENZ's start in the late 1990s")
    parser.add_argument("--country-source", default=COUNTRY_SOURCE_FILED,
                        choices=COUNTRY_SOURCES,
                        help="'filed' (default): every record is a New Zealand fire, the "
                             "points in the sea included; 'geometry': test the point "
                             "against the OCHA country polygons, leaving those out")
    output = parser.add_argument_group("output", "at least one is required")
    output.add_argument("--csv", type=Path, help="write the report to this .csv")
    output.add_argument("--docx", type=Path, help="write the report to this .docx (MS Word)")
    common.add_database_arguments(parser)
    parser.add_argument("--log-level", default=os.getenv("GISFIRE_LOG_LEVEL", "INFO"),
                        choices=["CRITICAL", "ERROR", "WARNING", "INFO", "DEBUG"],
                        help="verbosity (env: GISFIRE_LOG_LEVEL, default INFO)")


def parse_arguments(argv: list[str] | None = None) -> argparse.Namespace:
    """Parse the command line."""
    parser = argparse.ArgumentParser(
        description="Burnt-area statistics for the New Zealand Wildfire Occurrence "
                    "Database, with the records of each agency and the perimeters behind "
                    "each row.",
        epilog="Most WOD records are tiny incidents — the median FENZ area is five "
               "square metres — so use --min-area to ask about fires of a size. Database "
               "settings not given here are read from the environment (.env).",
    )
    parser.add_argument("--area-method", default=AREA_REPORTED, choices=AREA_METHODS,
                        help="'reported' (default): the area the agency recorded, every "
                             "fire; 'perimeter': the measured area of the published "
                             "perimeter, only the fires that have one")
    parser.add_argument("--min-area", type=hectares, default=None, metavar="HECTARES",
                        help="count only fires of at least this many hectares")
    parser.add_argument("--cause", default=None, choices=CAUSES,
                        help="restrict to one cause class")
    add_common_arguments(parser)
    arguments = parser.parse_args(argv)
    if arguments.csv is None and arguments.docx is None:
        parser.error("nothing to write: pass --csv, --docx, or both")
    return arguments


def compute(session: Session, year: int | None, logger: logging.Logger,
            financial_year: bool = False, area_method: str = AREA_REPORTED,
            min_area: float | None = None,
            country_source: str = COUNTRY_SOURCE_FILED,
            cause: str | None = None, source: str | None = None) -> list[Row]:
    """Measure the fires a year at a time, returning the report's rows in order.

    Every statement runs in ``session``'s transaction, so against one snapshot.
    """
    if year is not None:
        years = [year]
    else:
        with common.Spinner("Finding the years the WOD covers", logger):
            years = list(session.scalars(years_query(financial_year, cause, source)))

    measured: list[Row] = []
    for index, measuring in enumerate(years, start=1):
        label = period_label(measuring, financial_year)
        with common.Spinner(f"Summing the {area_method} burnt area of the WOD fires "
                            f"({label}: {index} of {len(years)})", logger):
            measured += [
                Row(country=record.country, year=measuring,
                    minimum=float(record.minimum), maximum=float(record.maximum),
                    total=float(record.total), fires=record.fires, fenz=record.fenz,
                    doc=record.doc, perimeters=record.perimeters,
                    financial_year=financial_year)
                for record in session.execute(statistics_query(
                    measuring, financial_year, area_method, min_area, country_source,
                    cause, source))
            ]

    if country_source == COUNTRY_SOURCE_GEOMETRY:
        outside = session.scalar(outside_count(year, financial_year, area_method,
                                               min_area, cause, source)) or 0
        if outside:
            logger.warning("%d fire(s) have a point inside no country — most of them in "
                           "the sea, recorded at a wrong location — and are not counted; "
                           "--country-source filed counts them", outside)

    countries = ordered_countries(session, {row.country for row in measured})
    rows = summarise(measured, countries)
    logger.info("Computed %d rows over %d year(s) (%s area, %s, country from %s)",
                len(rows), len({row.year for row in measured}), area_method,
                "fire years" if financial_year else "calendar years", country_source)
    return rows


def write_csv(rows: list[Row], path: Path, logger: logging.Logger) -> None:
    """Write the report as CSV: four decimals, because most areas are tiny."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(COLUMNS)
        for row in rows:
            writer.writerow(row.values)
    logger.info("Wrote %s", path)


def write_table(document, columns: tuple[str, ...], rows: list) -> None:
    """Add the report's table to a Word document, summary rows in bold."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for cell, heading in zip(table.rows[0].cells, columns):
        cell.text = heading
        cell.paragraphs[0].runs[0].bold = True
    for row in rows:
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, row.readable_values)):
            cell.text = value
            paragraph = cell.paragraphs[0]
            if index >= FIRST_NUMERIC_COLUMN:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for run in paragraph.runs:
                run.bold = row.is_total
                run.font.size = Pt(8)


def coverage_paragraph() -> str:
    """The coverage caveat every WOD report opens with."""
    return ("Coverage is not uniform. Before 1998 the WOD holds DOC records only and, in "
            "the words of its authors, does not give good overall coverage; from 1998 to "
            "2011 it holds FENZ and DOC records, deduplicated by the compilers; from 2012 "
            "FENZ records only. Most records are small incidents: the median FENZ area is "
            "five square metres.")


def write_docx(rows: list[Row], path: Path, logger: logging.Logger, year: int | None,
               financial_year: bool, area_method: str, min_area: float | None,
               country_source: str, cause: str | None, source: str | None) -> None:
    """Write the report as a Word document, its caveats first."""
    from docx import Document

    document = Document()
    document.add_heading("WOD wildfire burnt area (New Zealand)", level=1)
    scope = [f"year: {period_label(year, financial_year)}" if year is not None
             else "all years"]
    scope.append("fire years (1 July to 30 June)" if financial_year else "calendar years")
    if cause is not None:
        scope.append(f"only {cause} fires")
    if source is not None:
        scope.append(f"only {source} records")
    if min_area is not None:
        scope.append(f"only fires of {min_area:g} ha or more")
    if country_source == COUNTRY_SOURCE_GEOMETRY:
        scope.append("only fires whose point falls inside a country")
    if area_method == AREA_REPORTED:
        area = ("Areas in hectares as the agencies recorded them (AreaHa); a recorded 0 is "
                "counted.")
    else:
        area = ("Areas in hectares measured from the published perimeters on NZTM2000; "
                "only the fires with a perimeter are counted, and there are none for "
                "2012-2016.")
    document.add_paragraph(f"{area} Scope: {'; '.join(scope)}.")
    document.add_paragraph(coverage_paragraph())
    write_table(document, COLUMNS, rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    logger.info("Wrote %s", path)


def report(args: argparse.Namespace, engine: Engine, logger: logging.Logger) -> list[Row]:
    """Compute the statistics and write whichever outputs were asked for."""
    with Session(engine) as session:
        rows = compute(session, args.year, logger, args.financial_year, args.area_method,
                       args.min_area, args.country_source, args.cause, args.source)
    if not rows:
        raise RuntimeError(
            "No wildfires matched. Check --year and the other filters, and that the WOD "
            f"is imported — the import reads from {new_zealand_wod.FIRST_YEAR} on unless "
            "run with --all-years.")
    if args.csv is not None:
        write_csv(rows, args.csv, logger)
    if args.docx is not None:
        write_docx(rows, args.docx, logger, args.year, args.financial_year,
                   args.area_method, args.min_area, args.country_source, args.cause,
                   args.source)
    return rows


def run_application(argv: list[str] | None, parse, run, name: str) -> int:
    """The command-line boundary both reports share: settings, engine, errors."""
    args = parse(argv)
    logging.basicConfig(level=args.log_level,
                        format="%(asctime)s %(levelname)s %(message)s")
    logger = logging.getLogger(name)
    try:
        settings = common.resolve_database_settings(args)
    except RuntimeError as error:
        logger.error("%s", error)
        return 1
    engine = create_engine(common.database_url(settings))
    try:
        run(args, engine, logger)
    except Exception as error:  # noqa: BLE001  (the CLI boundary: report, do not traceback)
        logger.error("%s", error)
        return 1
    finally:
        engine.dispose()
    return 0


def main(argv: list[str] | None = None) -> int:
    return run_application(argv, parse_arguments, report, "wod-statistics")


if __name__ == "__main__":  # pragma nocover
    sys.exit(main())
