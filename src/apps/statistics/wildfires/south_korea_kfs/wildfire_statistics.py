#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Burnt-area statistics for the Korea Forest Service wildfire statistics.

Reports the release per year: how many fires there were, the smallest single fire,
the largest single fire and the total burnt area in hectares::

    python3 -m src.apps.statistics.wildfires.south_korea_kfs.wildfire_statistics \\
        --csv burnt.csv --docx burnt.docx

    python3 -m src.apps.statistics.wildfires.south_korea_kfs.wildfire_statistics \\
        --province 경북 --csv gyeongbuk.csv

Its six columns are the shared ones of every burnt-area report, in their order, so
the CSVs can be concatenated with the others.

The hectares are reported, not measured
----------------------------------------

The service publishes no perimeter, so there is no ``--area-method``: the figure is
``피해면적_합계``, the one burnt area each row publishes, stored as
:attr:`~src.providers.south_korea_kfs.wildfire.KfsWildfire.area_ha`. Unlike the Greek
statistic there is only one area per fire, so there is no ``--surface`` either.

Which year a fire counts towards
---------------------------------

:attr:`~src.providers.south_korea_kfs.wildfire.KfsWildfire.year`, the published year
the fire started in — the year the service totals it under.

**The last year of a release is partial**: the 2025-09-12 release ends on 11
September 2025. The report says so when a year is incomplete, by the date of its last
fire, rather than leaving a reader to compare 397 fires with the 756 of 2022.

There is no country to choose
-----------------------------

No ``--country`` and no ``--country-source``: the service publishes the fires of
South Korea and nothing else, and there is no point to test against a boundary in any
case. The ``Country`` column is the constant :data:`COUNTRY_NAME`. ``--province``
selects one province by its published abbreviation instead (``경북``, ``강원``…).

The area is dominated by single fires
-------------------------------------

The two Uiseong fires of 22 March 2025 are 99,282 ha, three quarters of everything in
the release, and are filed at the township they started in although they burnt into
four other municipalities. A total by province is a total by **ignition** province.
"""

from __future__ import annotations

import argparse
import csv
import datetime
import logging
import math
import os
import sys

from dataclasses import dataclass
from pathlib import Path

from sqlalchemy import Engine
from sqlalchemy import Select
from sqlalchemy import create_engine
from sqlalchemy import func
from sqlalchemy import literal
from sqlalchemy import select
from sqlalchemy.orm import Session

import src.settings  # noqa: F401  (imported for the side effect of loading .env)

from src.apps.imports import common
from src.data_model.wildfire import Wildfire
from src.providers import south_korea_kfs
from src.providers.south_korea_kfs.wildfire import KfsWildfire

#: Label used in the ``Year`` column for the summary row.
TOTAL_LABEL = "Total"

#: The report's columns: the six every burnt-area report shares, in their order.
COLUMNS = ("Country", "Year", "Fires", "Minimum (ha)", "Maximum (ha)", "Total (ha)")

#: Index of the first column that holds a number, right-aligned in the Word table.
FIRST_NUMERIC_COLUMN = 2

#: The country every fire in this dataset is in, and the whole ``Country`` column.
COUNTRY_NAME = "South Korea"

#: The ``kfs_wildfire`` table, joined by table so SQLAlchemy adds no polymorphic
#: join of its own.
KFS = KfsWildfire.__table__


def scope_conditions(year: int | None = None, province: str | None = None) -> list:
    """The ``WHERE`` conditions selecting the fires in scope."""
    conditions = []
    if year is not None:
        conditions.append(KFS.c.year == year)
    if province is not None:
        conditions.append(KFS.c.province_name == province)
    return conditions


def statistics_query(year: int | None = None, min_area: float | None = None,
                     province: str | None = None) -> Select:
    """Build the statistics query: one row per year, newest first.

    Returns
    -------
    Select
        ``country, year, minimum, maximum, total, fires``. The summary row is
        :func:`summarise`'s work.

    Notes
    -----
    ``min_area`` selects the fires the figures are computed from, before the
    aggregates; it does not discard years whose total came out small.
    """
    statistics = (
        select(
            literal(COUNTRY_NAME).label("country"),
            KFS.c.year.label("year"),
            func.min(KFS.c.area_ha).label("minimum"),
            func.max(KFS.c.area_ha).label("maximum"),
            func.sum(KFS.c.area_ha).label("total"),
            func.count().label("fires"),
        )
        .select_from(Wildfire)
        .join(KFS, KFS.c.id == Wildfire.id)
        .group_by(KFS.c.year)
        .order_by(KFS.c.year.desc())
    )
    for condition in scope_conditions(year, province):
        statistics = statistics.where(condition)
    if min_area is not None:
        statistics = statistics.where(KFS.c.area_ha >= min_area)
    return statistics


def last_fire_query(year: int) -> Select:
    """The local date of the last fire stored for ``year``, whatever the scope."""
    local = func.timezone(Wildfire.time_zone, Wildfire.start_date_time)
    return (
        select(func.max(func.date(local)))
        .select_from(Wildfire)
        .join(KFS, KFS.c.id == Wildfire.id)
        .where(KFS.c.year == year)
    )


def partial_years(session: Session, years: list[int]) -> dict[int, datetime.date]:
    """The years whose last fire is before December, with that fire's date.

    A year with fires into December is taken as complete. One whose last fire is
    earlier is almost certainly the release's cut-off — South Korea has fires every
    month — and is reported as partial.
    """
    partial = {}
    for year in years:
        last = session.scalar(last_fire_query(year))
        if last is not None and last.month < 12:
            partial[year] = last
    return partial


@dataclass(frozen=True)
class Row:
    """One line of the report.

    Attributes
    ----------
    country : str
        Always :data:`COUNTRY_NAME`.
    year : int or None
        The published year, or ``None`` for the summary row.
    minimum, maximum, total : float
        Smallest single fire, largest single fire and sum of every fire, in hectares
        as published.
    fires : int
        How many fires the three figures were computed from.
    """

    country: str
    year: int | None
    minimum: float
    maximum: float
    total: float
    fires: int

    @property
    def is_total(self) -> bool:
        """Whether this is the summary row rather than one of the years."""
        return self.year is None

    @property
    def year_label(self) -> str:
        return TOTAL_LABEL if self.is_total else str(self.year)

    @property
    def values(self) -> tuple[str, ...]:
        """The row as the CSV writes it, in :data:`COLUMNS` order."""
        return (self.country, self.year_label, str(self.fires),
                f"{self.minimum:.2f}", f"{self.maximum:.2f}", f"{self.total:.2f}")

    @property
    def readable_values(self) -> tuple[str, ...]:
        """The row as the Word document writes it: the numbers with separators."""
        return (self.country, self.year_label, f"{self.fires:,}",
                f"{self.minimum:,.2f}", f"{self.maximum:,.2f}", f"{self.total:,.2f}")


def combine(rows: list[Row], country: str = COUNTRY_NAME, year: int | None = None) -> Row:
    """One row summarising several: a minimum of minima, a sum of sums, and so on."""
    return Row(
        country=country,
        year=year,
        minimum=min(row.minimum for row in rows),
        maximum=max(row.maximum for row in rows),
        total=math.fsum(row.total for row in rows),
        fires=sum(row.fires for row in rows),
    )


def summarise(measured: list[Row]) -> list[Row]:
    """The years newest first, then the total; empty if nothing was measured."""
    if not measured:
        return []
    rows = sorted(measured, key=lambda row: row.year, reverse=True)
    return rows + [combine(rows)]


def hectares(text: str) -> float:
    """Argparse type for ``--min-area``: a finite, non-negative number of hectares."""
    try:
        area = float(text)
    except ValueError:
        raise argparse.ArgumentTypeError(f"{text!r} is not a number of hectares")
    if not math.isfinite(area):
        raise argparse.ArgumentTypeError(f"{text!r} is not a finite number of hectares")
    if area < 0:
        raise argparse.ArgumentTypeError(
            f"a burnt area cannot be negative, and {area:g} ha is: pass 0 or more, or "
            f"leave --min-area out to count every fire"
        )
    return area


def province(text: str) -> str:
    """Argparse type for ``--province``: a published province abbreviation."""
    name = text.strip()
    if name not in south_korea_kfs.PROVINCES:
        raise argparse.ArgumentTypeError(
            f"{text!r} is not a published province; use one of "
            f"{', '.join(south_korea_kfs.PROVINCES)}"
        )
    return name


def add_scope_arguments(parser: argparse.ArgumentParser) -> None:
    """The options selecting the fires, shared with the causes report."""
    parser.add_argument("-y", "--year", type=int,
                        help="restrict to one year, e.g. 2023: the year the fire started")
    parser.add_argument("--province", type=province,
                        help="restrict to one province, by its published abbreviation "
                             "(경북, 강원, …); fires are filed where they started")
    # Accepted only so that they can be refused with a reason: anyone reaching for one
    # has copied a command line from another report.
    parser.add_argument("--country", help=argparse.SUPPRESS)
    parser.add_argument("--country-source", help=argparse.SUPPRESS)
    parser.add_argument("--area-method", help=argparse.SUPPRESS)


def add_output_arguments(parser: argparse.ArgumentParser) -> None:
    """``--csv``, ``--docx``, the database settings and ``--log-level``."""
    output = parser.add_argument_group("output", "at least one is required")
    output.add_argument("--csv", type=Path, help="write the report to this .csv")
    output.add_argument("--docx", type=Path, help="write the report to this .docx (MS Word)")
    common.add_database_arguments(parser)
    parser.add_argument("--log-level", default=os.getenv("GISFIRE_LOG_LEVEL", "INFO"),
                        choices=["CRITICAL", "ERROR", "WARNING", "INFO", "DEBUG"],
                        help="verbosity (env: GISFIRE_LOG_LEVEL, default INFO)")


def check_arguments(parser: argparse.ArgumentParser, arguments: argparse.Namespace,
                    outputs: tuple[str, ...] = ("csv", "docx")) -> None:
    """Refuse the options this dataset has no use for, saying why.

    ``outputs`` names the output options of which at least one is required.
    """
    if arguments.country is not None or arguments.country_source is not None:
        parser.error(
            f"there is no --country or --country-source here: the Korea Forest Service "
            f"publishes the fires of {COUNTRY_NAME} and nothing else, with no coordinate "
            f"to test against a boundary. Use --province to select part of the country."
        )
    if arguments.area_method is not None:
        parser.error(
            "there is no --area-method here: this dataset publishes no perimeter, so "
            "nothing is measured. The hectares are the ones the service reports."
        )
    if all(getattr(arguments, name) is None for name in outputs):
        options = [f"--{name.replace('_', '-')}" for name in outputs]
        parser.error(f"nothing to write: pass {', '.join(options[:-1])} or {options[-1]}, "
                     f"or several")


def parse_arguments(argv: list[str] | None = None) -> argparse.Namespace:
    """Parse the command line."""
    parser = argparse.ArgumentParser(
        description="Burnt-area statistics for the Korea Forest Service wildfire "
                    "statistics (산불통계데이터).",
        epilog="Areas are the hectares the service reports — this dataset publishes no "
               "perimeter. Every fire is South Korean, so there is no --country. "
               "Database settings not given here are read from the environment (.env).",
    )
    add_scope_arguments(parser)
    parser.add_argument("--min-area", type=hectares, default=None, metavar="HECTARES",
                        help="count only fires of at least this many hectares; by default "
                             "every fire counts")
    add_output_arguments(parser)
    arguments = parser.parse_args(argv)
    check_arguments(parser, arguments)
    return arguments


def scope_description(year: int | None, province_name: str | None,
                      min_area: float | None = None) -> str:
    """The scope in words, for the Word document."""
    scope = [f"year: {year}" if year is not None else "all years"]
    if province_name is not None:
        scope.append(f"province: {province_name} "
                     f"({south_korea_kfs.PROVINCES[province_name]})")
    if min_area is not None:
        scope.append(f"only fires of {min_area:g} ha or more")
    return "; ".join(scope)


def partial_paragraph(partial: dict[int, datetime.date]) -> str | None:
    """The sentence saying which years are incomplete, or ``None``."""
    if not partial:
        return None
    years = ", ".join(f"{year} (last fire {last.isoformat()})"
                      for year, last in sorted(partial.items()))
    return (f"Incomplete year(s): {years}. A release ends part way through its last "
            f"year, so that year's figures are not comparable with full years.")


def compute(session: Session, year: int | None, logger: logging.Logger,
            min_area: float | None = None,
            province_name: str | None = None) -> tuple[list[Row], dict[int, datetime.date]]:
    """Run the statement; return the report's rows in order and the partial years."""
    with common.Spinner(f"Summing the burnt area of the {COUNTRY_NAME} fires", logger):
        measured = [
            Row(country=record.country, year=record.year,
                minimum=float(record.minimum), maximum=float(record.maximum),
                total=float(record.total), fires=record.fires)
            for record in session.execute(statistics_query(year, min_area, province_name))
        ]
        partial = partial_years(session, [row.year for row in measured])

    rows = summarise(measured)
    logger.info("Computed %d rows over %d year(s) (%s)", len(rows), len(measured),
                "every fire" if min_area is None else f"fires of {min_area:g} ha or more")
    paragraph = partial_paragraph(partial)
    if paragraph:
        logger.warning("%s", paragraph)
    return rows, partial


def write_csv(rows: list, path: Path, logger: logging.Logger,
              columns: tuple[str, ...] = COLUMNS) -> None:
    """Write a report as CSV: bare numbers, no thousands separators."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(columns)
        for row in rows:
            writer.writerow(row.values)
    logger.info("Wrote %s", path)


def write_table(document, columns: tuple[str, ...], rows: list,
                first_numeric: int = FIRST_NUMERIC_COLUMN) -> None:
    """Add one table to a Word document, summary rows in bold."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for cell, heading in zip(table.rows[0].cells, columns):
        cell.text = heading
        cell.paragraphs[0].runs[0].bold = True
    for row in rows:
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, row.readable_values)):
            cell.text = value
            paragraph = cell.paragraphs[0]
            if index >= first_numeric:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for run in paragraph.runs:
                run.bold = row.is_total
                run.font.size = Pt(9)


def write_docx(rows: list[Row], path: Path, logger: logging.Logger,
               year: int | None = None, min_area: float | None = None,
               province_name: str | None = None,
               partial: dict[int, datetime.date] | None = None) -> None:
    """Write the report as a Word document, its caveats first."""
    # Imported here so that --csv keeps working without python-docx.
    from docx import Document

    document = Document()
    document.add_heading(f"Korea Forest Service wildfire burnt area ({COUNTRY_NAME})",
                         level=1)
    document.add_paragraph(
        f"Areas in hectares as the service reports them (피해면적_합계) — this dataset "
        f"publishes no perimeter, so nothing is measured. Years are the year the fire "
        f"started. Scope: {scope_description(year, province_name, min_area)}."
    )
    document.add_paragraph(
        "Fires are filed at the place they started: a fire that spread into other "
        "municipalities or provinces counts wholly towards the one it started in."
    )
    paragraph = partial_paragraph(partial or {})
    if paragraph:
        document.add_paragraph(paragraph)
    write_table(document, COLUMNS, rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    logger.info("Wrote %s", path)


def no_fires_message(min_area: float | None = None) -> str:
    """Why a report is empty, for the error that refuses to write it."""
    threshold = "" if min_area is None else \
        f" No fire reached the --min-area of {min_area:g} ha."
    return (f"No wildfires matched. Check --year and --province, and that the "
            f"{COUNTRY_NAME} fires are imported with "
            f"src.apps.imports.wildfires.south_korea_kfs.import_wildfires.{threshold}")


def report(args: argparse.Namespace, engine: Engine, logger: logging.Logger) -> list[Row]:
    """Compute the statistics and write whichever outputs were asked for."""
    with Session(engine) as session:
        rows, partial = compute(session, args.year, logger, args.min_area, args.province)
    if not rows:
        raise RuntimeError(no_fires_message(args.min_area))
    if args.csv is not None:
        write_csv(rows, args.csv, logger)
    if args.docx is not None:
        write_docx(rows, args.docx, logger, args.year, args.min_area, args.province,
                   partial)
    return rows


def run_application(argv: list[str] | None, parse, run, logger_name: str) -> int:
    """Parse, connect, run and report errors without a traceback — the CLI boundary."""
    args = parse(argv)
    logging.basicConfig(level=args.log_level, format="%(asctime)s %(levelname)s %(message)s")
    logger = logging.getLogger(logger_name)

    try:
        settings = common.resolve_database_settings(args)
    except RuntimeError as error:
        logger.error("%s", error)
        return 1

    engine = create_engine(common.database_url(settings))
    try:
        run(args, engine, logger)
    except Exception as error:  # noqa: BLE001  (the CLI boundary: report, do not traceback)
        logger.error("%s", error)
        return 1
    finally:
        engine.dispose()
    return 0


def main(argv: list[str] | None = None) -> int:
    return run_application(argv, parse_arguments, report, "south-korea-kfs-statistics")


if __name__ == "__main__":  # pragma nocover
    sys.exit(main())
