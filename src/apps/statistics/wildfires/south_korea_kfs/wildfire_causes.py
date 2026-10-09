#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Wildfire counts by cause for the Korea Forest Service wildfire statistics.

Two tables over the same fires and the same scope as
:mod:`~src.apps.statistics.wildfires.south_korea_kfs.wildfire_statistics`, so the
``Country``, ``Year`` and ``Fires`` columns agree with it row for row::

    python3 -m src.apps.statistics.wildfires.south_korea_kfs.wildfire_causes \\
        --csv causes.csv --classes-csv classes.csv --docx causes.docx

**The shared fire causes** (``--csv``), in the columns of the New Zealand report
(:mod:`~src.apps.statistics.wildfires.new_zealand_wod.wildfire_causes`), so the two
countries can be set side by side::

    Country      Year   Fires  Determined  Lightning  Lightning (%)  Natural  Human  Unknown  Lightning (ha)  Lightning (% of ha)

**The eighteen cause classes** (``--classes-csv``), this dataset's own: for each year
and class the fires, their share of the year's fires, the burnt area and its share,
and how many of the fires give the cause as suspected rather than established::

    Country      Year  Class    Fires  Fires (%)  Area (ha)  Area (%)  Suspected

The Word document (``--docx``) holds both. At least one output is required.

Where the classes come from
----------------------------

The published coded cause is *other (free text)* on three rows in four, so both
tables classify the free text, ``발생원인_기타``: the import translates each of its
spellings and assigns one of
:data:`~src.providers.south_korea_kfs.CAUSE_CLASSES`, and reduces that to one of the
four :data:`~src.providers.south_korea_kfs.FIRE_CAUSES`. Lightning is the
``lightning`` class and nothing else; ``natural`` is always zero, no Korean text
naming another natural cause; *unknown* and *rekindled from an earlier fire* are
``unknown``.

A **suspected** cause (``추정``, ``의심``, ``수사중``) is classified by what it
suspects — one fire in five. The ``Suspected`` column says how many, so a reader can
see how much of a class rests on suspicion.

Why the denominators differ
---------------------------

In the shared table the percentages are of the **determined** fires, as in every
cause report, with ``Unknown`` shown so that what was left out is visible. In the
class table they are of **all** the year's fires and all its area, ``unknown`` being
one of the classes: the classes add up to 100%.

A percentage of nothing is an empty cell, not zero.
"""

from __future__ import annotations

import argparse
import logging
import math
import sys

from dataclasses import dataclass
from pathlib import Path

from sqlalchemy import Engine
from sqlalchemy import Select
from sqlalchemy import func
from sqlalchemy import select
from sqlalchemy.orm import Session

from src.apps.statistics.wildfires.south_korea_kfs.wildfire_statistics import COUNTRY_NAME
from src.apps.statistics.wildfires.south_korea_kfs.wildfire_statistics import KFS
from src.apps.statistics.wildfires.south_korea_kfs.wildfire_statistics import TOTAL_LABEL
from src.apps.statistics.wildfires.south_korea_kfs.wildfire_statistics import (
    add_output_arguments,
)
from src.apps.statistics.wildfires.south_korea_kfs.wildfire_statistics import (
    add_scope_arguments,
)
from src.apps.statistics.wildfires.south_korea_kfs.wildfire_statistics import (
    check_arguments,
)
from src.apps.statistics.wildfires.south_korea_kfs.wildfire_statistics import (
    no_fires_message,
)
from src.apps.statistics.wildfires.south_korea_kfs.wildfire_statistics import (
    partial_paragraph,
)
from src.apps.statistics.wildfires.south_korea_kfs.wildfire_statistics import (
    partial_years,
)
from src.apps.statistics.wildfires.south_korea_kfs.wildfire_statistics import (
    run_application,
)
from src.apps.statistics.wildfires.south_korea_kfs.wildfire_statistics import (
    scope_conditions,
)
from src.apps.statistics.wildfires.south_korea_kfs.wildfire_statistics import (
    scope_description,
)
from src.apps.statistics.wildfires.south_korea_kfs.wildfire_statistics import write_csv
from src.apps.statistics.wildfires.south_korea_kfs.wildfire_statistics import write_table
from src.data_model.wildfire import Wildfire
from src.providers import south_korea_kfs

#: The shared table's columns: the New Zealand cause report's, in its order.
COLUMNS = ("Country", "Year", "Fires", "Determined", "Lightning", "Lightning (%)",
           "Natural", "Human", "Unknown", "Lightning (ha)", "Lightning (% of ha)")

#: The class table's columns.
CLASS_COLUMNS = ("Country", "Year", "Class", "Fires", "Fires (%)", "Area (ha)",
                 "Area (%)", "Suspected")

#: The fire causes whose cause somebody determined: the shared table's denominator.
DETERMINED_CAUSES = (south_korea_kfs.CAUSE_LIGHTNING, south_korea_kfs.CAUSE_NATURAL,
                     south_korea_kfs.CAUSE_HUMAN)


def percentage(part: float, whole: float) -> str:
    """``part`` as a percentage of ``whole``, written out; empty where there is no whole."""
    return "" if not whole else f"{100.0 * part / whole:.2f}"


# --------------------------------------------------------------------------
# The shared fire causes
# --------------------------------------------------------------------------

def _from_kfs(statement: Select, year: int | None, province_name: str | None) -> Select:
    statement = statement.select_from(Wildfire).join(KFS, KFS.c.id == Wildfire.id)
    for condition in scope_conditions(year, province_name):
        statement = statement.where(condition)
    return statement


def counts_query(year: int | None = None, province_name: str | None = None) -> Select:
    """One row per year: ``year, fires, lightning, natural, human, unknown,
    lightning_hectares, determined_hectares``."""
    cause = KFS.c.fire_cause
    area = KFS.c.area_ha
    lightning = cause == south_korea_kfs.CAUSE_LIGHTNING
    statement = select(
        KFS.c.year.label("year"),
        func.count().label("fires"),
        func.count().filter(lightning).label("lightning"),
        func.count().filter(cause == south_korea_kfs.CAUSE_NATURAL).label("natural"),
        func.count().filter(cause == south_korea_kfs.CAUSE_HUMAN).label("human"),
        func.count().filter(cause == south_korea_kfs.CAUSE_UNKNOWN).label("unknown"),
        func.coalesce(func.sum(area).filter(lightning), 0.0).label("lightning_hectares"),
        func.coalesce(func.sum(area).filter(cause.in_(DETERMINED_CAUSES)),
                      0.0).label("determined_hectares"),
    )
    return (_from_kfs(statement, year, province_name)
            .group_by(KFS.c.year).order_by(KFS.c.year.desc()))


@dataclass(frozen=True)
class Row:
    """One line of the shared table.

    ``fires`` is the sum of the four cause counts; ``year`` is ``None`` for the total.
    """

    country: str
    year: int | None
    fires: int
    lightning: int
    natural: int
    human: int
    unknown: int
    lightning_hectares: float
    determined_hectares: float

    @property
    def is_total(self) -> bool:
        return self.year is None

    @property
    def year_label(self) -> str:
        return TOTAL_LABEL if self.is_total else str(self.year)

    @property
    def determined(self) -> int:
        """The fires with a known cause: the denominator."""
        return self.lightning + self.natural + self.human

    @property
    def values(self) -> tuple[str, ...]:
        """The row as the CSV writes it, in :data:`COLUMNS` order."""
        return (self.country, self.year_label, str(self.fires), str(self.determined),
                str(self.lightning), percentage(self.lightning, self.determined),
                str(self.natural), str(self.human), str(self.unknown),
                f"{self.lightning_hectares:.4f}",
                percentage(self.lightning_hectares, self.determined_hectares))

    @property
    def readable_values(self) -> tuple[str, ...]:
        """The row as the Word document writes it: numbers with separators."""
        return (self.country, self.year_label, f"{self.fires:,}", f"{self.determined:,}",
                f"{self.lightning:,}", percentage(self.lightning, self.determined),
                f"{self.natural:,}", f"{self.human:,}", f"{self.unknown:,}",
                f"{self.lightning_hectares:,.2f}",
                percentage(self.lightning_hectares, self.determined_hectares))


def combine(rows: list[Row]) -> Row:
    """The total row: counts and hectares added, shares recomputed from them."""
    return Row(
        country=COUNTRY_NAME, year=None,
        fires=sum(row.fires for row in rows),
        lightning=sum(row.lightning for row in rows),
        natural=sum(row.natural for row in rows),
        human=sum(row.human for row in rows),
        unknown=sum(row.unknown for row in rows),
        lightning_hectares=math.fsum(row.lightning_hectares for row in rows),
        determined_hectares=math.fsum(row.determined_hectares for row in rows),
    )


def compute(session: Session, year: int | None = None,
            province_name: str | None = None) -> list[Row]:
    """The shared table's rows: the years newest first, then the total."""
    measured = [
        Row(country=COUNTRY_NAME, year=record.year, fires=record.fires,
            lightning=record.lightning, natural=record.natural, human=record.human,
            unknown=record.unknown,
            lightning_hectares=float(record.lightning_hectares),
            determined_hectares=float(record.determined_hectares))
        for record in session.execute(counts_query(year, province_name))
    ]
    return measured + [combine(measured)] if measured else []


# --------------------------------------------------------------------------
# The eighteen classes
# --------------------------------------------------------------------------

def classes_query(year: int | None = None, province_name: str | None = None) -> Select:
    """One row per year and class: ``year, cause_class, fires, hectares, suspected``."""
    statement = select(
        KFS.c.year.label("year"),
        KFS.c.cause_class.label("cause_class"),
        func.count().label("fires"),
        func.sum(KFS.c.area_ha).label("hectares"),
        func.count().filter(KFS.c.cause_suspected).label("suspected"),
    )
    return (_from_kfs(statement, year, province_name)
            .group_by(KFS.c.year, KFS.c.cause_class))


@dataclass(frozen=True)
class ClassRow:
    """One line of the class table.

    ``year_fires`` and ``year_hectares`` are the whole year's, the denominators of the
    two shares; ``year`` is ``None`` for the rows over every year.
    """

    country: str
    year: int | None
    cause_class: str
    fires: int
    hectares: float
    suspected: int
    year_fires: int
    year_hectares: float

    @property
    def is_total(self) -> bool:
        return self.year is None

    @property
    def year_label(self) -> str:
        return TOTAL_LABEL if self.is_total else str(self.year)

    @property
    def values(self) -> tuple[str, ...]:
        """The row as the CSV writes it, in :data:`CLASS_COLUMNS` order."""
        return (self.country, self.year_label, self.cause_class, str(self.fires),
                percentage(self.fires, self.year_fires), f"{self.hectares:.2f}",
                percentage(self.hectares, self.year_hectares), str(self.suspected))

    @property
    def readable_values(self) -> tuple[str, ...]:
        """The row as the Word document writes it: numbers with separators."""
        return (self.country, self.year_label, self.cause_class, f"{self.fires:,}",
                percentage(self.fires, self.year_fires), f"{self.hectares:,.2f}",
                percentage(self.hectares, self.year_hectares), f"{self.suspected:,}")


def _class_block(year: int | None, counts: dict[str, tuple[int, float, int]]) -> list[ClassRow]:
    """The rows of one year (or of all years), most fires first, ties by class order."""
    year_fires = sum(fires for fires, _, _ in counts.values())
    year_hectares = math.fsum(hectares for _, hectares, _ in counts.values())
    order = sorted(counts, key=lambda name: (-counts[name][0],
                                              south_korea_kfs.CAUSE_CLASSES.index(name)))
    return [ClassRow(country=COUNTRY_NAME, year=year, cause_class=name,
                     fires=counts[name][0], hectares=counts[name][1],
                     suspected=counts[name][2], year_fires=year_fires,
                     year_hectares=year_hectares)
            for name in order]


def compute_classes(session: Session, year: int | None = None,
                    province_name: str | None = None) -> list[ClassRow]:
    """The class table's rows: each year newest first, then every year together.

    Only the classes that occur are listed; a class with no fire in a year has no row
    for it.
    """
    by_year: dict[int, dict[str, tuple[int, float, int]]] = {}
    overall: dict[str, list] = {}
    for record in session.execute(classes_query(year, province_name)):
        hectares = float(record.hectares)
        by_year.setdefault(record.year, {})[record.cause_class] = (
            record.fires, hectares, record.suspected)
        total = overall.setdefault(record.cause_class, [0, [], 0])
        total[0] += record.fires
        total[1].append(hectares)
        total[2] += record.suspected
    if not by_year:
        return []
    rows: list[ClassRow] = []
    for counting in sorted(by_year, reverse=True):
        rows += _class_block(counting, by_year[counting])
    rows += _class_block(None, {name: (fires, math.fsum(areas), suspected)
                                for name, (fires, areas, suspected) in overall.items()})
    return rows


# --------------------------------------------------------------------------
# The application
# --------------------------------------------------------------------------

def parse_arguments(argv: list[str] | None = None) -> argparse.Namespace:
    """Parse the command line."""
    parser = argparse.ArgumentParser(
        description="Wildfire counts by cause for the Korea Forest Service wildfire "
                    "statistics: the four shared fire causes and the eighteen classes "
                    "of the published free-text cause.",
        epilog="The shared table's percentages are of the fires whose cause is known; "
               "the class table's are of all the fires. Database settings not given "
               "here are read from the environment (.env).",
    )
    add_scope_arguments(parser)
    parser.add_argument("--classes-csv", type=Path,
                        help="write the eighteen-class table to this .csv")
    add_output_arguments(parser)
    arguments = parser.parse_args(argv)
    check_arguments(parser, arguments, outputs=("csv", "classes_csv", "docx"))
    return arguments


def write_docx(rows: list[Row], class_rows: list[ClassRow], path: Path,
               logger: logging.Logger, year: int | None = None,
               province_name: str | None = None, partial: dict | None = None) -> None:
    """Write both tables as a Word document, their caveats first."""
    from docx import Document

    document = Document()
    document.add_heading(f"Korea Forest Service wildfires by cause ({COUNTRY_NAME})",
                         level=1)
    document.add_paragraph(
        "Causes are classified from the published free-text cause (발생원인_기타), "
        "translated and grouped by GisFIRE; the published coded cause is 'other (free "
        "text)' on most rows. A cause given as suspected or under investigation is "
        f"classified by what it suspects. Scope: {scope_description(year, province_name)}."
    )
    paragraph = partial_paragraph(partial or {})
    if paragraph:
        document.add_paragraph(paragraph)

    document.add_heading("Shared fire causes", level=2)
    document.add_paragraph(
        "Lightning is the lightning class alone; no Korean cause text names another "
        "natural cause. Unknown includes fires rekindled from an earlier one. The "
        "percentages are of the fires whose cause is known; hectares are as published."
    )
    write_table(document, COLUMNS, rows)

    document.add_heading("Cause classes", level=2)
    document.add_paragraph(
        "The percentages are of all the year's fires and burnt area, so a year's "
        "classes add up to 100%. Suspected counts the fires whose cause is suspected "
        "rather than established. Single large fires dominate the area: the two "
        "Uiseong fires of March 2025 are agricultural burning and grave visitor."
    )
    write_table(document, CLASS_COLUMNS, class_rows, first_numeric=3)

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    logger.info("Wrote %s", path)


def report(args: argparse.Namespace, engine: Engine,
           logger: logging.Logger) -> tuple[list[Row], list[ClassRow]]:
    """Count the causes and write whichever outputs were asked for."""
    with Session(engine) as session:
        rows = compute(session, args.year, args.province)
        class_rows = compute_classes(session, args.year, args.province)
        partial = partial_years(session, [row.year for row in rows if not row.is_total])
    if not rows:
        raise RuntimeError(no_fires_message())

    total = rows[-1]
    logger.info("%d of %d fire(s) have a known cause; %d are lightning (%s%%), "
                "accounting for %s%% of the determined fires' area",
                total.determined, total.fires, total.lightning,
                percentage(total.lightning, total.determined),
                percentage(total.lightning_hectares, total.determined_hectares))
    paragraph = partial_paragraph(partial)
    if paragraph:
        logger.warning("%s", paragraph)

    if args.csv is not None:
        write_csv(rows, args.csv, logger, COLUMNS)
    if args.classes_csv is not None:
        write_csv(class_rows, args.classes_csv, logger, CLASS_COLUMNS)
    if args.docx is not None:
        write_docx(rows, class_rows, args.docx, logger, args.year, args.province, partial)
    return rows, class_rows


def main(argv: list[str] | None = None) -> int:
    return run_application(argv, parse_arguments, report, "south-korea-kfs-causes")


if __name__ == "__main__":  # pragma nocover
    sys.exit(main())
